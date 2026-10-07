from pathlib import Path

import pytest
from openpyxl import Workbook

from excel_twin import normalise


def build(path: Path, python_values: dict | None = None, break_formula: bool = False) -> int:
    """A tiny model: inputs, a calculation, and a Reconciliation sheet in the default layout."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Model"
    ws["A1"], ws["B1"] = "Price", 22.83
    ws["A2"], ws["B2"] = "Quantity", 1.5
    ws["A3"], ws["B3"] = "Revenue", "=ROUND(B1*B2,2)"          # 34.245 -> 34.25, Excel's tie
    ws["A4"], ws["B4"] = "Label", '=IF(B3>30,"high","low")'
    ws["A5"], ws["B5"] = "Ratio", "=B3/0" if break_formula else "=B3/B2"
    expected = {"Revenue": 34.25, "Label": "high", "Ratio": 34.25 / 1.5}
    expected.update(python_values or {})
    r = wb.create_sheet("Reconciliation")
    rows = [("Model", "Revenue", expected["Revenue"], "Model!B3"), ("Model", "Label", expected["Label"], "Model!B4"),
            ("Model", "Ratio", expected["Ratio"], "Model!B5")]
    first, last = 7, 6 + len(rows)
    r["B3"] = (f'=IF(COUNTIF(G{first}:G{last},"No")=0,"All "&COUNTA(B{first}:B{last})&" checks match",'
               f'COUNTIF(G{first}:G{last},"No")&" checks do not match")')
    for i, (area, item, value, ref) in enumerate(rows, start=first):
        r[f"A{i}"], r[f"B{i}"], r[f"D{i}"], r[f"E{i}"] = area, item, value, f"={ref}"
        r[f"G{i}"] = (f'=IF(AND(D{i}="",E{i}=""),"Yes",IF(AND(ISNUMBER(D{i}),ISNUMBER(E{i})),'
                      f'IF(ABS(E{i}-D{i})<=1E-9*MAX(1,ABS(D{i})),"Yes","No"),IF(D{i}=E{i},"Yes","No")))')
    wb.save(path)
    normalise(path)
    return len(rows)


@pytest.fixture
def make(tmp_path):
    def _make(name="book.xlsx", **kw):
        p = tmp_path / name
        n = build(p, **kw)
        return p, n
    return _make
